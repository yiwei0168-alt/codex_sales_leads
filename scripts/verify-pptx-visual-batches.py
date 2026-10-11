"""Local synthetic boundary probe with actual RapidOCR, not content-quality acceptance."""
from pathlib import Path
from io import BytesIO
from hashlib import sha256
import json
import subprocess
import sys
import tempfile
from PIL import Image, ImageDraw
from pptx import Presentation

root = Path.cwd().resolve()
with tempfile.TemporaryDirectory(dir=root / "tmp") as temp:
    folder = Path(temp); source = folder / "batch.pptx"; artifact = folder / "artifact.json"
    deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6]); blocks = []
    for i in range(3):
        image = Image.new("RGB", (180, 60), "white")
        ImageDraw.Draw(image).text((8, 18), f"LOCAL BATCH {i + 1}", fill="black")
        stream = BytesIO(); image.save(stream, format="PNG"); stream.seek(0)
        shape = slide.shapes.add_picture(stream, 0, 0, 1800000, 600000)
        blocks.append(dict(id=f"image-{i}", unitIndex=1, unitType="slide", quality="pending-ocr",
                           shapePath=[i+1], shapeIdPath=[shape.shape_id]))
    deck.save(source)
    artifact.write_text(json.dumps(dict(sourceSha256=sha256(source.read_bytes()).hexdigest(), blocks=blocks)), encoding="utf-8")
    output = folder / "output"
    subprocess.run([sys.executable, str(root / "scripts/prepare_pptx_visual_review.py"), "--input", str(source),
                    "--artifact", str(artifact), "--output", str(output), "--max-unique-images", "2", "--all-batches"],
                   check=True, timeout=300)
    result = json.loads((output / "review.json").read_text(encoding="utf-8"))
    assert result["imageBatches"] == 2 and result["uniqueImagesProcessed"] == 3
    assert result["qualityAccepted"] is False and result["extractionStatusChanged"] is False
    assert len(result["records"]) == 3 and result["externalCalls"] == 0
    for record in result["records"]:
        assert record["reason"] != "local-image-budget-exhausted"
        assert sha256((output / record["imageFile"]).read_bytes()).hexdigest() == record["reviewImageSha256"]
    evidence = dict(syntheticBoundaryProbe=True, actualLocalRapidOCR=True, objects=3, batches=2, batchLimit=2,
                    allImagesHashVerified=True, qualityAccepted=False, externalCalls=0, productionDeployed=False)
    (root / "tmp/pptx-visual-batches-result.json").write_text(json.dumps(evidence, indent=2), encoding="utf-8")
    print(json.dumps(evidence))
