from io import BytesIO
from types import SimpleNamespace
import unittest
import xml.etree.ElementTree as ET
from pathlib import Path
from hashlib import sha256
import json
import os
import tempfile
from unittest.mock import patch
import numpy  # Load once before patch.dict restores the module registry between runs.
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import MSO_CONNECTOR
from prepare_pptx_visual_review import connector_candidate, diagram_candidate, resolve_shape, visible_image, main, checkpoint_lines


class VisualReviewTest(unittest.TestCase):
    def test_checkpoint_under_realistic_long_owner_job_path(self):
        with tempfile.TemporaryDirectory(dir=Path.cwd() / "tmp") as temp:
            directory=Path(temp)/('a'*60)/('b'*60)/('c'*40)
            directory.mkdir(parents=True)
            lines=[{"text":"candidate","confidence":.9,"box":[[0,0]]}]
            self.assertEqual(checkpoint_lines(directory,'d'*64,lambda:lines),(lines,False))
            self.assertEqual(checkpoint_lines(directory,'d'*64,lambda:[]),(lines,True))
            stored=directory/('d'*64+'.json')
            if os.name=='nt': stored=Path('\\\\?\\'+str(stored.resolve()))
            stored.unlink()

    def test_checkpoint_corruption_and_interrupted_atomic_write(self):
        with tempfile.TemporaryDirectory(dir=Path.cwd() / "tmp") as temp:
            folder = Path(temp); image_hash = "a" * 64
            lines = [{"text": "unreviewed", "confidence": .8, "box": [[0, 0]]}]
            self.assertEqual(checkpoint_lines(folder, image_hash, lambda: lines), (lines, False))
            def forbidden(): raise AssertionError("Must reuse complete checkpoint")
            self.assertEqual(checkpoint_lines(folder, image_hash, forbidden), (lines, True))
            file = folder / f"{image_hash}.json"
            file.write_text('{"body":', encoding="utf-8")
            self.assertEqual(checkpoint_lines(folder, image_hash, lambda: lines), (lines, False))
            stored = json.loads(file.read_text(encoding="utf-8")); stored["body"]["lines"][0]["text"] = "changed"
            file.write_text(json.dumps(stored), encoding="utf-8")
            with patch("prepare_pptx_visual_review.os.replace", side_effect=RuntimeError("interrupted")):
                with self.assertRaises(RuntimeError): checkpoint_lines(folder, image_hash, lambda: lines)
            self.assertEqual(list(folder.glob("*.tmp")), [])
            self.assertEqual(checkpoint_lines(folder, image_hash, lambda: lines), (lines, False))

    def test_all_batches_cover_101_distinct_images_without_accepting_quality(self):
        # Real PPTX serialization/crops/package output; OCR is a deterministic local stub.
        with tempfile.TemporaryDirectory(dir=Path.cwd() / "tmp") as temp:
            folder = Path(temp); source = folder / "many.pptx"; artifact = folder / "artifact.json"
            deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6]); blocks = []
            for i in range(101):
                stream = BytesIO(); Image.new("RGB", (4, 4), (i, 0, 0)).save(stream, format="PNG"); stream.seek(0)
                shape = slide.shapes.add_picture(stream, 0, 0, 100, 100)
                blocks.append(dict(id=f"image-{i}", unitIndex=1, unitType="slide", quality="pending-ocr", shapePath=[i+1], shapeIdPath=[shape.shape_id]))
            deck.save(source)
            artifact.write_text(json.dumps(dict(sourceSha256=sha256(source.read_bytes()).hexdigest(), blocks=blocks)), encoding="utf-8")
            calls = []
            def engine(image, **_kwargs):
                calls.append(int(image[0, 0, 0])); return SimpleNamespace(txts=None, scores=None, boxes=None)
            module = SimpleNamespace(RapidOCR=lambda **_kwargs: engine, EngineType=SimpleNamespace(ONNXRUNTIME=1), LangRec=SimpleNamespace(CH=1))
            for all_batches in (False, True):
                calls.clear(); output = folder / str(all_batches)
                argv = ["prepare", "--input", str(source), "--artifact", str(artifact), "--output", str(output), "--max-unique-images", "100"]
                if all_batches: argv.append("--all-batches")
                with patch("sys.argv", argv), patch.dict("sys.modules", {"rapidocr": module}), patch("socket.socket.connect"), patch("socket.create_connection"):
                    main()
                result = json.loads((output / "review.json").read_text(encoding="utf-8"))
                self.assertEqual(len(result["records"]), 101)
                self.assertEqual(len(calls), 101 if all_batches else 100)
                self.assertEqual(result["imageBatches"], 2 if all_batches else 1)
                self.assertFalse(result["qualityAccepted"])
                self.assertFalse(result["extractionStatusChanged"])
                self.assertEqual(result["externalCalls"], 0)
                exhausted = [r for r in result["records"] if r["reason"] == "local-image-budget-exhausted"]
                self.assertEqual(len(exhausted), 0 if all_batches else 1)
                for record in result["records"]:
                    if "imageFile" in record:
                        self.assertEqual(sha256((output / record["imageFile"]).read_bytes()).hexdigest(), record["reviewImageSha256"])

    def test_connector_keeps_native_geometry_without_inventing_attachments(self):
        deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 100, 200, 400, 600)
        result = connector_candidate(line)
        self.assertEqual(result["offset"], {"x": "100", "y": "200"})
        self.assertEqual(result["extent"], {"cx": "300", "cy": "400"})
        self.assertEqual(result["coordinateSpace"], "parent-group-emu")
        self.assertIsNone(result["startAttachment"])
        self.assertIsNone(result["endAttachment"])
        target = slide.shapes.add_textbox(0, 0, Inches(1), Inches(1))
        line.begin_connect(target, 0)
        attached = connector_candidate(line)
        self.assertEqual(attached["startAttachment"]["id"], str(target.shape_id))
        self.assertIsNone(connector_candidate(target))

    def test_native_diagram_binding_and_external_rejection(self):
        element = ET.fromstring('''<shape xmlns:d="http://schemas.openxmlformats.org/drawingml/2006/diagram"
          xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><d:relIds r:dm="data1"/></shape>''')
        blob = b'''<d:dataModel xmlns:d="http://schemas.openxmlformats.org/drawingml/2006/diagram"
          xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><d:ptLst>
          <d:pt modelId="node-a"><d:t><a:p><a:r><a:t>Label</a:t></a:r></a:p></d:t></d:pt>
          </d:ptLst><d:cxnLst><d:cxn srcId="node-a" destId="node-b"/></d:cxnLst></d:dataModel>'''
        relation = SimpleNamespace(is_external=False, target_part=SimpleNamespace(blob=blob))
        shape = SimpleNamespace(element=element, part=SimpleNamespace(rels={"data1": relation}))
        result = diagram_candidate(shape)
        self.assertEqual(result["points"][0]["text"], "Label")
        self.assertEqual(result["connections"][0]["srcId"], "node-a")
        relation.is_external = True
        with self.assertRaises(ValueError): diagram_candidate(shape)

    def test_source_crop_excludes_hidden_pixels(self):
        image = Image.new("RGB", (100, 50), "white")
        stream = BytesIO(); image.save(stream, format="PNG")
        shape = SimpleNamespace(image=SimpleNamespace(blob=stream.getvalue()),crop_left=.2,crop_top=0,
                                crop_right=.1,crop_bottom=.2,rotation=90)
        cropped, info = visible_image(shape)
        self.assertEqual(cropped.size, (70, 40))
        self.assertEqual(info["coordinateSpace"], "cropped-image-pixels")
        self.assertEqual(info["rotation"], 90)
        shape.crop_left = -.1
        with self.assertRaises(ValueError): visible_image(shape)

    def test_nested_native_id_must_match(self):
        deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
        group = slide.shapes.add_group_shape()
        child = group.shapes.add_textbox(0, 0, Inches(1), Inches(1))
        block = {"unitIndex": 1, "shapePath": [1, 1], "shapeIdPath": [group.shape_id, child.shape_id]}
        self.assertEqual(resolve_shape(deck, block).shape_id, child.shape_id)
        block["shapeIdPath"] = [group.shape_id, child.shape_id + 1]
        with self.assertRaises(ValueError): resolve_shape(deck, block)
        block["shapePath"] = [0, 1]
        with self.assertRaises(ValueError): resolve_shape(deck, block)


if __name__ == "__main__":
    unittest.main()
