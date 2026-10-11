"""Local public-source fixture for version replacement; never edits the original."""
from pathlib import Path
import sys
import pymupdf
from pptx import Presentation
from pptx.util import Inches
from openpyxl import load_workbook

source, destination = Path(sys.argv[1]), Path(sys.argv[2]).resolve()
root = (Path.cwd() / 'tmp' / 'binary-revisions').resolve()
if root not in destination.parents:
    raise SystemExit('Fixture destination outside dedicated local directory')
marker = sys.argv[3]
destination.parent.mkdir(parents=True, exist_ok=True)
if source.suffix.lower() == '.pdf':
    doc = pymupdf.open(source)
    page = doc.new_page()
    page.insert_text((50, 80), marker)
    doc.save(destination)
    doc.close()
elif source.suffix.lower() == '.pptx':
    deck = Presentation(source)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1)).text = marker
    deck.save(destination)
elif source.suffix.lower() == '.xlsx':
    book = load_workbook(source)
    sheet = book.create_sheet('Revision verification')
    sheet.append(['Revision marker'])
    sheet.append([marker])
    book.save(destination)
else:
    raise SystemExit('Unsupported fixture format')
