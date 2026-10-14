"""Regression checks for original spreadsheet values and physical row positions."""
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as SheetImage
from openpyxl.worksheet.table import Table
import pymupdf
from pptx import Presentation
from pptx.util import Inches
from PIL import Image
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from knowledge_extraction_v2 import extract_file


class SpreadsheetCoordinatesTest(unittest.TestCase):
    def test_unsupported_table_style_metadata_preserves_cells_and_merges(self):
        with TemporaryDirectory() as directory:
            original=Path(directory)/"original.xlsx"
            patched=Path(directory)/"table-style.xlsx"
            book=Workbook(); sheet=book.active
            sheet.append(["Area","Rate"]);sheet.append(["Orkney",30]);sheet.append(["Shetland",20])
            sheet.merge_cells("D1:E1");sheet["D1"]="Source note"
            sheet.add_table(Table(displayName="Rates",ref="A1:B3"))
            book.save(original);book.close()
            with zipfile.ZipFile(original) as source, zipfile.ZipFile(patched,"w") as output:
                for part in source.infolist():
                    data=source.read(part.filename)
                    if part.filename=="xl/tables/table1.xml":
                        root=ET.fromstring(data)
                        namespace="{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
                        root.find(f"{namespace}tableColumns/{namespace}tableColumn").set("dxfId","0")
                        data=ET.tostring(root,encoding="utf-8")
                    output.writestr(part,data)
            artifact=extract_file(patched,"table-style.xlsx")
            self.assertEqual(artifact["qualitySummary"],{"success":1,"blank":0,"pending-ocr":0,"failed":0})
            table=artifact["blocks"][0]
            self.assertIn("D1:E1",table["metadata"]["mergedRanges"])
            self.assertEqual(table["table"]["rows"][1][:2],["Shetland","20"])
            self.assertEqual(table["table"]["cellRange"],"A1:E3")

    def test_cached_formula_error_table_requires_review(self):
        with TemporaryDirectory() as directory:
            path=Path(directory)/"cached-errors.xlsx"
            book=Workbook(); report=book.active; report.title="Report"
            report["A1"]="Metric"; report["B2"]="#VALUE!"; report["B2"].data_type="e"
            literal=book.create_sheet("Literal"); literal["A1"]="The formula may show #VALUE!"
            book.save(path); book.close()
            artifact=extract_file(path,"cached-errors.xlsx")
            self.assertEqual(artifact["qualitySummary"],{"success":1,"blank":0,"pending-ocr":1,"failed":0})
            table=artifact["blocks"][0]
            self.assertEqual((table["id"],table["quality"]),("sheet-1-table","pending-ocr"))
            self.assertEqual(table["metadata"]["reviewReason"],"cached-spreadsheet-error")
            self.assertEqual(table["metadata"]["cachedErrorCells"],[{"cell":"B2","value":"#VALUE!"}])
            self.assertEqual(artifact["blocks"][1]["quality"],"success")

    def test_mixed_and_image_only_sheets_are_not_complete(self):
        with TemporaryDirectory() as directory:
            path=Path(directory)/"visual.xlsx"; picture=Path(directory)/"image.png"
            Image.new("RGB",(20,20),"black").save(picture)
            book=Workbook(); sheet=book.active; sheet.title="Chart"
            sheet.append(["Value"]); sheet.append([0]); sheet.append([5])
            chart=BarChart(); chart.add_data(Reference(sheet,min_col=1,min_row=1,max_row=3),titles_from_data=True)
            sheet.add_chart(chart,"D2")
            book.create_sheet("Image").add_image(SheetImage(str(picture)),"B3")
            chart_page=book.create_chartsheet("ChartPage"); chart2=BarChart()
            chart2.add_data(Reference(sheet,min_col=1,min_row=1,max_row=3),titles_from_data=True); chart_page.add_chart(chart2)
            book.create_sheet("Blank"); book.save(path); book.close()
            artifact=extract_file(path,"visual.xlsx")
            self.assertEqual(artifact["qualitySummary"],{"success":0,"blank":1,"pending-ocr":3,"failed":0})
            pending=[b for b in artifact["blocks"] if b["quality"]=="pending-ocr"]
            self.assertEqual([b["section"] for b in pending],["Chart","Image","ChartPage"])
            self.assertEqual(pending[0]["metadata"]["sourceReference"]["kind"],"drawing")
            self.assertEqual(pending[0]["metadata"]["sourceReference"]["worksheetPart"],"xl/worksheets/sheet1.xml")
            self.assertEqual(artifact["blocks"][0]["table"]["rows"][0],["0"])

    def test_blank_rows_zero_and_false_keep_original_positions(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "coordinates.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Original"
            sheet["B2"] = "Value"
            sheet["B4"] = 0
            sheet["C4"] = False
            sheet["B6"] = -1.5
            workbook.create_sheet("Blank")
            workbook.save(path)
            workbook.close()
            artifact = extract_file(path, "coordinates.xlsx")
            self.assertEqual(artifact["qualitySummary"]["success"], 1)
            self.assertEqual(artifact["qualitySummary"]["blank"], 1)
            block = artifact["blocks"][0]
            self.assertEqual(block["section"], "Original")
            self.assertEqual(block["unitIndex"], 1)
            table = block["table"]
            self.assertEqual(table["startRow"], 2)
            self.assertEqual(table["cellRange"], "A1:C6")
            self.assertEqual(table["headerRow"], 1)
            self.assertEqual(table["startColumn"], 1)
            self.assertEqual(table["headers"], ["", "", ""])
            physical_rows = [table["headers"], *table["rows"]]
            self.assertEqual(len(physical_rows), 6)
            self.assertEqual(physical_rows[1][1], "Value")
            self.assertEqual(physical_rows[2], ["", "", ""])
            self.assertEqual(physical_rows[3][1:], ["0", "False"])
            self.assertEqual(physical_rows[4], ["", "", ""])
            self.assertEqual(physical_rows[5][1], "-1.5")

    def test_zero_only_sheet_is_not_blank(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "zero.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = 0
            workbook.save(path)
            workbook.close()
            artifact = extract_file(path, "zero.xlsx")
            self.assertEqual(artifact["qualitySummary"]["success"], 1)
            self.assertEqual(artifact["blocks"][0]["table"]["headers"], ["0"])


class PptxGroupTest(unittest.TestCase):
    def test_visual_only_and_mixed_slides_require_review(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "visuals.pptx"
            picture = Path(directory) / "image.png"
            Image.new("RGB", (20, 20), "black").save(picture)
            deck = Presentation()
            deck.slides.add_slide(deck.slide_layouts[6])
            image_slide = deck.slides.add_slide(deck.slide_layouts[6])
            image_slide.shapes.add_picture(str(picture), 0, 0)
            mixed = deck.slides.add_slide(deck.slide_layouts[6])
            mixed.shapes.add_textbox(0, 0, Inches(1), Inches(1)).text = "Caption"
            group = mixed.shapes.add_group_shape()
            group.shapes.add_picture(str(picture), 0, 0)
            chart_slide = deck.slides.add_slide(deck.slide_layouts[6])
            data = CategoryChartData()
            data.categories = ["A", "B"]
            data.add_series("Values", [0, 5])
            chart_slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 0, 0, Inches(3), Inches(3), data)
            deck.save(path)
            artifact = extract_file(path, "visuals.pptx")
            self.assertEqual(artifact["qualitySummary"], {
                "success": 0, "blank": 1, "pending-ocr": 3, "failed": 0})
            pending = [b for b in artifact["blocks"] if b["quality"] == "pending-ocr"]
            self.assertEqual([b["unitIndex"] for b in pending], [2, 3, 4])
            self.assertEqual(pending[1]["shapePath"], [2, 1])
            self.assertTrue(any(b["text"] == "Caption" for b in artifact["blocks"]))

    def test_nested_text_has_unique_order_and_native_object_paths(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "groups.pptx"
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            direct = slide.shapes.add_textbox(0, 0, Inches(1), Inches(1))
            direct.text = "Direct"
            group = slide.shapes.add_group_shape()
            child = group.shapes.add_textbox(0, 0, Inches(1), Inches(1))
            child.text = "Grouped"
            nested = group.shapes.add_group_shape()
            inner = nested.shapes.add_textbox(0, 0, Inches(1), Inches(1))
            inner.text = "Deep"
            deck.save(path)
            artifact = extract_file(path, "groups.pptx")
            blocks = artifact["blocks"]
            self.assertEqual([b["text"] for b in blocks], ["Direct", "Grouped", "Deep"])
            self.assertEqual([b["shapePath"] for b in blocks], [[1], [2, 1], [2, 2, 1]])
            self.assertEqual(blocks[2]["shapeIdPath"], [group.shape_id, nested.shape_id, inner.shape_id])
            self.assertEqual(len(set(b["id"] for b in blocks)), 3)


class PdfBlankPageTest(unittest.TestCase):
    def test_text_does_not_hide_visual_content(self):
        with TemporaryDirectory() as directory:
            path=Path(directory)/"mixed.pdf"; picture=Path(directory)/"image.png"
            Image.new("RGB",(20,20),"black").save(picture)
            document=pymupdf.open(); page=document.new_page(); page.insert_text((50,50),"Caption")
            page.insert_image(pymupdf.Rect(50,80,100,130),filename=str(picture))
            page=document.new_page(); page.insert_text((50,50),"Chart caption")
            page.draw_rect(pymupdf.Rect(50,80,100,130),color=(0,0,0))
            document.save(path); document.close()
            artifact=extract_file(path,"mixed.pdf")
            self.assertEqual(artifact["qualitySummary"]["pending-ocr"],2)
            self.assertTrue(any(b["text"]=="Caption" for b in artifact["blocks"]))
            pending=[b for b in artifact["blocks"] if b["quality"]=="pending-ocr"]
            self.assertEqual([b["unitIndex"] for b in pending],[1,2])
            self.assertTrue(all(len(b["bbox"])==4 for b in pending))

    def test_only_structurally_empty_pages_are_blank(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "pages.pdf"
            document = pymupdf.open()
            document.new_page()  # No page content.
            document.new_page().insert_text((50, 50), "Original text")
            page = document.new_page()
            page.draw_rect(pymupdf.Rect(10, 10, 100, 100), color=(0, 0, 0))
            document.new_page().add_text_annot((50, 50), "Visible annotation")
            document.save(path)
            document.close()
            artifact = extract_file(path, "pages.pdf")
            self.assertEqual(artifact["qualitySummary"], {
                "success": 1, "blank": 1, "pending-ocr": 2, "failed": 0})
            self.assertEqual(artifact["blocks"][0]["quality"], "blank")
            self.assertEqual(artifact["blocks"][0]["unitIndex"], 1)


if __name__ == "__main__":
    unittest.main()
