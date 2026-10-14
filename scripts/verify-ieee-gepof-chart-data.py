"""Read the embedded MSGraph data for the 2014 GEPOF slide 16 review."""

import argparse
import hashlib
import io
import json
import struct

import olefile
from pptx import Presentation


SOURCE_SHA256 = "030a730584498f8b98109a1bf2da22a29cc443021c63fbd1c18cf2bad2da6910"
EXPECTED_CATEGORIES = ["10 mm", "8 mm", "6 mm", "5 mm"]
EXPECTED_SERIES = {
    "Bending at RX": [0.82, 1.46, 2.22, 2.68],
    "Bending at TX": [0.56, 0.97, 1.95, 2.67],
}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--input",
        default="tmp/ma24-independent-quarantine/ieee-gepof-tutorial-v14.pptx",
    )
    args = parser.parse_args()
    with open(args.input, "rb") as stream:
        source = stream.read()
    assert hashlib.sha256(source).hexdigest() == SOURCE_SHA256

    presentation = Presentation(io.BytesIO(source))
    assert len(presentation.slides) == 82
    shape = presentation.slides[15].shapes[3]
    assert shape.shape_id == 6 and shape.ole_format.prog_id == "MSGraph.Chart.8"
    with olefile.OleFileIO(io.BytesIO(shape.ole_format.blob)) as embedded:
        workbook = embedded.openstream("Workbook").read()

    labels: dict[tuple[int, int], str] = {}
    numbers: dict[tuple[int, int], float] = {}
    axis_labels: list[str] = []
    position = 0
    while position + 4 <= len(workbook):
        kind, length = struct.unpack_from("<HH", workbook, position)
        end = position + 4 + length
        assert end <= len(workbook), "Truncated MSGraph record"
        data = workbook[position + 4 : end]
        if kind == 0x0004:  # LABEL
            assert len(data) >= 9
            row, column = struct.unpack_from("<HH", data)
            labels[row, column] = data[9:].decode("utf-16le")
        elif kind == 0x0003:  # NUMBER
            assert len(data) == 15
            row, column = struct.unpack_from("<HH", data)
            numbers[row, column] = struct.unpack_from("<d", data, 7)[0]
        elif kind == 0x100D and "Attenuation".encode("utf-16le") in data:
            axis_labels.append(data[4:].decode("utf-16le"))
        position = end
    assert position == len(workbook)
    assert [labels[0, column] for column in range(1, 5)] == EXPECTED_CATEGORIES
    assert {key for key in numbers} == {(row, column) for row in (1, 2) for column in range(1, 5)}
    assert len(axis_labels) == 1 and axis_labels[0] == "Attenuation (dBo)"

    series: dict[str, list[float]] = {}
    for row in (1, 2):
        name = labels[row, 0]
        values = [round(numbers[row, column], 2) for column in range(1, 5)]
        assert values == EXPECTED_SERIES[name]
        series[name] = values
    assert series == EXPECTED_SERIES
    print(
        json.dumps(
            {
                "sourceSha256": SOURCE_SHA256,
                "slide": 16,
                "shapeId": shape.shape_id,
                "oleProgId": shape.ole_format.prog_id,
                "workbookSha256": hashlib.sha256(workbook).hexdigest(),
                "axisLabelLiteral": axis_labels[0],
                "categories": EXPECTED_CATEGORIES,
                "series": series,
                "readOnly": True,
            }
        )
    )


if __name__ == "__main__":
    main()
