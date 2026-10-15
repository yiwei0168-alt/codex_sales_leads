"""Bounded lexical overlap screen for the current private G14 source originals."""

import hashlib
import json
import re
import unicodedata
from collections import defaultdict
from itertools import combinations
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parent.parent
INPUT = ROOT / "tmp/ma24-independent-quarantine/g14-source-binary-closure-v9-2026-10-16.json"
OUTPUT = ROOT / "tmp/ma24-independent-quarantine/g14-cross-source-text-overlap-v9-2026-10-16.json"
EXPECTED_POOL_SHA = "96459f096b17bafdd32dfedc18db924e0339ae53dd4b563ea1d2d7f78b88b9ed"
WORD = re.compile(r"[a-z0-9]+", re.ASCII)
SHINGLE_WIDTHS = (8, 12)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for part in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(part)
    return digest.hexdigest()


def tokens(text: str) -> list[str]:
    return WORD.findall(unicodedata.normalize("NFKC", text).lower())


def pdf_units(path: Path) -> tuple[list[str], int]:
    reader = PdfReader(str(path))
    return [(page.extract_text() or "") for page in reader.pages], len(reader.pages)


def xlsx_units(path: Path) -> tuple[list[str], int]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        units = []
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                cells = [str(value) for value in row if value is not None]
                if cells:
                    rows.append(" ".join(cells))
            units.append("\n".join(rows))
        return units, len(workbook.worksheets)
    finally:
        workbook.close()


def pptx_units(source_hash: str) -> tuple[list[str], int, str]:
    folder = ROOT / "tmp/ma24-independent-quarantine"
    candidates = []
    for path in set(folder.glob("*extraction*.json")) | set(folder.glob("*product-extraction*.json")):
        data = json.loads(path.read_text(encoding="utf-8"))
        documents = data.get("documents", [data]) if isinstance(data, dict) else []
        for document in documents:
            if document.get("sourceSha256") == source_hash:
                version = document.get("extractorVersion", data.get("extractorVersion", ""))
                version_key = tuple(int(value) for value in re.findall(r"\d+", version))
                candidates.append((version_key, str(path), document))
    assert candidates, f"No source-hash-bound PPTX extraction: {source_hash}"
    _, selected, document = max(candidates, key=lambda item: (item[0], item[1]))
    slides = defaultdict(list)
    for block in document.get("blocks", []):
        if isinstance(block.get("text"), str) and block["text"].strip():
            slides[block.get("unitIndex")].append(block["text"])
    assert slides, f"Empty PPTX extraction: {source_hash}"
    return ["\n".join(slides[index]) for index in sorted(slides)], len(slides), selected


closure = json.loads(INPUT.read_text(encoding="utf-8"))
assert closure["poolSha256"] == EXPECTED_POOL_SHA
assert sha256(ROOT / "tmp/ma24-independent-quarantine/g14-current-draft-pool-v9-metadata-2026-10-16.json") == EXPECTED_POOL_SHA
assert len(closure["originals"]) == closure["distinctOriginalBinaries"] == 16

documents = []
shingles = {}
for entry in closure["originals"]:
    path = (ROOT / entry["path"]).resolve()
    assert path.is_relative_to(ROOT / "tmp/ma24-independent-quarantine") or path.is_relative_to(ROOT / "tmp/ma24-independent-candidates")
    assert sha256(path) == entry["sha256"], entry["path"]
    if path.suffix.lower() == ".pdf":
        units, unit_count = pdf_units(path)
        extraction = "pypdf text"
    elif path.suffix.lower() == ".xlsx":
        units, unit_count = xlsx_units(path)
        extraction = "openpyxl cached values"
    elif path.suffix.lower() == ".pptx":
        units, unit_count, extraction_path = pptx_units(entry["sha256"])
        extraction = f"existing product extraction: {Path(extraction_path).name}"
    else:
        raise AssertionError(f"Unsupported original format: {path.suffix}")
    unique = {width: set() for width in SHINGLE_WIDTHS}
    words = 0
    for unit in units:
        sequence = tokens(unit)
        words += len(sequence)
        for width in SHINGLE_WIDTHS:
            unique[width].update(tuple(sequence[index:index + width])
                                 for index in range(max(0, len(sequence) - width + 1)))
    assert words > 0, entry["path"]
    shingles[entry["sha256"]] = unique
    documents.append(dict(sha256=entry["sha256"], path=entry["path"],
                          families=entry["families"], format=path.suffix.lower(),
                          extraction=extraction, units=unit_count,
                          tokens=words,
                          uniqueShingles={str(width): len(unique[width]) for width in SHINGLE_WIDTHS}))

cross_family_pairs = []
same_family_pairs = []
for left, right in combinations(documents, 2):
    item = dict(leftSha256=left["sha256"], rightSha256=right["sha256"],
                leftFamily=left["families"], rightFamily=right["families"])
    for width in SHINGLE_WIDTHS:
        shared = shingles[left["sha256"]][width] & shingles[right["sha256"]][width]
        item[f"shared{width}TokenShingles"] = len(shared)
        item[f"smaller{width}TokenShare"] = round(len(shared) / max(1, min(
            left["uniqueShingles"][str(width)], right["uniqueShingles"][str(width)])), 6)
        item[f"sampleShared{width}TokenPhrases"] = [" ".join(words) for words in sorted(shared)[:3]]
    (same_family_pairs if set(left["families"]) & set(right["families"])
     else cross_family_pairs).append(item)
for pairs in (cross_family_pairs, same_family_pairs):
    pairs.sort(key=lambda item: (-item["shared8TokenShingles"],
                                 -item["shared12TokenShingles"], item["leftSha256"], item["rightSha256"]))
report = dict(schema="ma24-g14-cross-source-text-overlap-v1", inputSha256=sha256(INPUT),
              poolSha256=EXPECTED_POOL_SHA, shingleWidths=SHINGLE_WIDTHS,
              originals=documents, crossFamilyPairs=cross_family_pairs,
              sameFamilyPositiveControls=same_family_pairs,
              limitation="Lexical screen of local extractable text only. OCR gaps, charts, formulas, differently worded derivatives, unrecorded prompts and provider knowledge remain unassessed. Shared boilerplate is not evidence of answer leakage; zero shared shingles does not prove independence.",
              fullSetFrozen=False, goldApprovedByThisCheck=0, realDualArmRuns=0,
              g17ToG20Advanced=False)
OUTPUT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps(dict(report=str(OUTPUT.relative_to(ROOT)), reportSha256=sha256(OUTPUT),
                      originals=len(documents), crossFamilyPairs=len(cross_family_pairs),
                      crossPairsWithShared8=sum(item["shared8TokenShingles"] > 0 for item in cross_family_pairs),
                      crossPairsWithShared12=sum(item["shared12TokenShingles"] > 0 for item in cross_family_pairs),
                      sameFamilyPairs=len(same_family_pairs),
                      sameFamilyPairsWithShared12=sum(item["shared12TokenShingles"] > 0 for item in same_family_pairs),
                      topPairs=[dict(leftFamily=item["leftFamily"], rightFamily=item["rightFamily"],
                                     shared8TokenShingles=item["shared8TokenShingles"],
                                     shared12TokenShingles=item["shared12TokenShingles"],
                                     smaller8TokenShare=item["smaller8TokenShare"])
                                for item in cross_family_pairs[:5]],
                      positiveControls=[dict(leftFamily=item["leftFamily"],
                                             shared8TokenShingles=item["shared8TokenShingles"],
                                             shared12TokenShingles=item["shared12TokenShingles"])
                                        for item in same_family_pairs]), indent=2))
