"""Local OCR candidates for unresolved PPTX objects; never promotes extraction status."""
from pathlib import Path
from io import BytesIO
from hashlib import sha256
import argparse
import json
import socket
import os
import uuid
import xml.etree.ElementTree as ET

from PIL import Image
from pptx import Presentation


def checkpoint_lines(directory, image_hash, produce):
    """Checkpoints remain unreviewed candidates; corrupt/incomplete files are recomputed."""
    path = directory / f"{image_hash}.json" if directory else None
    if path:
        try:
            stored = json.loads(path.read_text(encoding="utf-8"))
            body = stored["body"]
            if (body["imageSha256"] == image_hash and isinstance(body["lines"], list)
                    and stored["sha256"] == sha256(json.dumps(body, sort_keys=True, allow_nan=False).encode()).hexdigest()
                    and all(isinstance(line["text"], str) and isinstance(line["confidence"], (int, float))
                            and isinstance(line["box"], list) for line in body["lines"])):
                return body["lines"], True
        except (OSError, ValueError, KeyError, TypeError):
            pass
    lines = produce()
    if path:
        body = {"imageSha256": image_hash, "lines": lines}
        data = json.dumps({"body": body, "sha256": sha256(json.dumps(body, sort_keys=True, allow_nan=False).encode()).hexdigest()}, allow_nan=False)
        temporary = path.with_suffix(f".{uuid.uuid4().hex}.tmp")
        try:
            with temporary.open("x", encoding="utf-8") as stream:
                stream.write(data); stream.flush(); os.fsync(stream.fileno())
            os.replace(temporary, path)
        except OSError as error:
            raise RuntimeError("Unable to persist OCR checkpoint") from error
        finally:
            temporary.unlink(missing_ok=True)
    return lines, False


def connector_candidate(shape):
    if shape.element.tag != "{http://schemas.openxmlformats.org/presentationml/2006/main}cxnSp":
        return None
    ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
          "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    def attrs(path):
        node = shape.element.find(path, ns)
        return dict(node.attrib) if node is not None else None
    return {"objectXmlSha256": sha256(ET.tostring(shape.element)).hexdigest(),
            "coordinateSpace": "parent-group-emu", "transform": attrs(".//a:xfrm"),
            "offset": attrs(".//a:xfrm/a:off"), "extent": attrs(".//a:xfrm/a:ext"),
            "geometry": attrs(".//a:prstGeom"), "startAttachment": attrs(".//a:stCxn"),
            "endAttachment": attrs(".//a:endCxn"), "headEnd": attrs(".//a:headEnd"),
            "tailEnd": attrs(".//a:tailEnd")}


def diagram_candidate(shape):
    ns = {"d": "http://schemas.openxmlformats.org/drawingml/2006/diagram",
          "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    rel_ids = shape.element.find(".//d:relIds", ns)
    if rel_ids is None:
        return None
    rel_id = rel_ids.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}dm")
    relation = shape.part.rels[rel_id]
    if relation.is_external:
        raise ValueError("External diagram data is not allowed")
    blob = relation.target_part.blob
    data = ET.fromstring(blob)
    points = [{"modelId": pt.get("modelId"), "type": pt.get("type"),
               "text": "\n".join("".join(p.itertext()) for p in pt.findall(".//a:p", ns)).strip()}
              for pt in data.findall(".//d:pt", ns)]
    connections = [{key: connection.get(key) for key in ("modelId", "type", "srcId", "destId", "srcOrd", "destOrd")}
                   for connection in data.findall(".//d:cxn", ns)]
    return {"diagramSha256": sha256(blob).hexdigest(), "points": points, "connections": connections}


def resolve_shape(deck, block):
    slide_index = block["unitIndex"]
    if not isinstance(slide_index, int) or not 1 <= slide_index <= len(deck.slides):
        raise ValueError("Invalid slide index")
    shapes = deck.slides[slide_index - 1].shapes
    path, ids = block["shapePath"], block["shapeIdPath"]
    if not path or len(path) != len(ids):
        raise ValueError("Invalid object path")
    for depth, (ordinal, expected_id) in enumerate(zip(path, ids)):
        if not isinstance(ordinal, int) or not 1 <= ordinal <= len(shapes):
            raise ValueError("Object path unavailable")
        shape = shapes[ordinal - 1]
        if shape.shape_id != expected_id:
            raise ValueError("Object identity changed")
        if depth + 1 < len(path):
            shapes = shape.shapes
    return shape


def visible_image(shape):
    blob = shape.image.blob
    image = Image.open(BytesIO(blob)).convert("RGB")
    crop = [float(getattr(shape, f"crop_{side}")) for side in ("left", "top", "right", "bottom")]
    if any(value < 0 or value >= 1 for value in crop) or crop[0]+crop[2] >= 1 or crop[1]+crop[3] >= 1:
        raise ValueError("Unsupported image crop")
    w, h = image.size
    image = image.crop((round(w*crop[0]), round(h*crop[1]), round(w*(1-crop[2])), round(h*(1-crop[3]))))
    if image.width < 1 or image.height < 1:
        raise ValueError("Empty image crop")
    # These are cropped source-image pixels, not transformed slide coordinates.
    return image, {"embeddedImageSha256": sha256(blob).hexdigest(), "crop": crop,
                   "rotation": float(shape.rotation), "coordinateSpace": "cropped-image-pixels"}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--artifact", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--max-unique-images", type=int, default=50)
    parser.add_argument("--all-batches", action="store_true", help="Process successive bounded image batches before publishing the package")
    parser.add_argument("--checkpoint-dir", help="Owned job-local checkpoint root for interrupted retries")
    parser.add_argument("--block-id", action="append", default=[], help="Process only selected unresolved object IDs")
    args = parser.parse_args()
    root = Path.cwd().resolve()
    source, artifact, output = (Path(value).resolve() for value in (args.input, args.artifact, args.output))
    if any(root not in path.parents for path in (source, artifact, output)) or source.suffix.lower() != ".pptx":
        raise ValueError("Local workspace PPTX, artifact and output required")
    if not 1 <= args.max_unique_images <= 100:
        raise ValueError("Image budget must be 1..100")
    if (output/"review.json").exists():
        raise ValueError("Completed review package exists; use a new output directory")
    payload = json.loads(artifact.read_text(encoding="utf-8"))
    if "documents" in payload:
        if not isinstance(payload["documents"], list) or len(payload["documents"]) != 1:
            raise ValueError("Exactly one uploaded document required")
        payload = payload["documents"][0]
    source_hash = sha256(source.read_bytes()).hexdigest()
    if payload["sourceSha256"] != source_hash:
        raise ValueError("Extraction source changed")
    pending = [b for b in payload["blocks"] if b["quality"] == "pending-ocr" and b["unitType"] == "slide"]
    total_pending = len(pending)
    if args.block_id:
        selected = set(args.block_id)
        if len(selected) != len(args.block_id) or not selected.issubset({b["id"] for b in pending}):
            raise ValueError("Selected object IDs must be unique unresolved blocks")
        pending = [b for b in pending if b["id"] in selected]
    deck = Presentation(source)
    # Validate every locator before starting any local model work.
    shapes = [resolve_shape(deck, b) for b in pending]
    output.mkdir(parents=True, exist_ok=True)
    models = root / ".models/docling-v3/RapidOcr"
    names = ["PP-OCRv6_det_small.onnx", "ch_ppocr_mobile_v2.0_cls_mobile.onnx", "PP-OCRv6_rec_small.onnx"]
    if not all((models/name).is_file() for name in names):
        raise ValueError("Existing local OCR models required; downloads disabled")
    model_hashes = {name: sha256((models/name).read_bytes()).hexdigest() for name in names}
    checkpoint = None
    if args.checkpoint_dir:
        checkpoint_root = Path(args.checkpoint_dir).resolve()
        if root not in checkpoint_root.parents:
            raise ValueError("Checkpoint directory must remain in workspace")
        identity = dict(source=source_hash, artifact=sha256(artifact.read_bytes()).hexdigest(), models=model_hashes,
                        implementation=sha256(Path(__file__).read_bytes()).hexdigest())
        checkpoint = checkpoint_root / sha256(json.dumps(identity, sort_keys=True).encode()).hexdigest()
        checkpoint.mkdir(parents=True, exist_ok=True)
    def deny_network(*_args, **_kwargs):
        raise RuntimeError("Network disabled during local visual review")
    socket.socket.connect = deny_network
    socket.create_connection = deny_network
    import numpy as np
    from rapidocr import RapidOCR, EngineType, LangRec
    engine = RapidOCR(params={
        "Det.model_path": models/names[0], "Cls.model_path": models/names[1], "Rec.model_path": models/names[2],
        "Det.engine_type": EngineType.ONNXRUNTIME, "Cls.engine_type": EngineType.ONNXRUNTIME,
        "Rec.engine_type": EngineType.ONNXRUNTIME, "Rec.lang_type": LangRec.CH, "Global.text_score": 0.35,
    })
    cache, records = {}, []
    processed_images, image_batches = 0, 0
    reused_images = 0
    for block, shape in zip(pending, shapes):
        record = {"blockId": block["id"], "unitIndex": block["unitIndex"], "shapePath": block["shapePath"],
                  "shapeIdPath": block["shapeIdPath"], "status": "review-required"}
        try:
            connector = connector_candidate(shape)
            if connector is not None:
                record["connector"] = connector
                record["reason"] = "native-connector-candidate-requires-review"
                records.append(record); continue
            diagram = diagram_candidate(shape)
            if diagram is not None:
                record["diagram"] = diagram
                record["reason"] = "native-diagram-candidate-requires-review"
                records.append(record); continue
            image, info = visible_image(shape)
            buf = BytesIO(); image.save(buf, format="PNG")
            data = buf.getvalue(); image_hash = sha256(data).hexdigest()
            record.update(info); record["reviewImageSha256"] = image_hash
            if image_hash not in cache:
                if len(cache) >= args.max_unique_images:
                    if not args.all_batches:
                        record["reason"] = "local-image-budget-exhausted"
                        records.append(record); continue
                    cache.clear()
                if not cache:
                    image_batches += 1
                (output/f"{image_hash}.png").write_bytes(data)
                def produce():
                    result = engine(np.asarray(image), text_score=0.35)
                    return [{"text": text, "confidence": float(score), "box": np.asarray(box).tolist()}
                            for text, score, box in zip(result.txts if result.txts is not None else [],
                                result.scores if result.scores is not None else [], result.boxes if result.boxes is not None else [])]
                lines, reused = checkpoint_lines(checkpoint, image_hash, produce)
                cache[image_hash] = lines
                processed_images += int(not reused)
                reused_images += int(reused)
            record["imageFile"] = f"{image_hash}.png"
            record["ocrLines"] = cache[image_hash]
            record["reason"] = "ocr-candidate-requires-visual-review" if cache[image_hash] else "no-text-image-still-requires-review"
        except (AttributeError, OSError, ValueError) as error:
            record["reason"] = "unsupported-visual-object"
            record["errorType"] = type(error).__name__
        records.append(record)
    review = {"schema": "pptx-visual-review-candidates-v1", "sourceSha256": source_hash,
              "extractionArtifactSha256": sha256(artifact.read_bytes()).hexdigest(),
              "modelSha256": model_hashes,
              "records": records, "totalUnresolvedObjects": total_pending, "selectedObjects": len(pending),
              "uniqueImagesProcessed": len({r["imageFile"] for r in records if "imageFile" in r}),
              "imageExecutions": processed_images, "checkpointReuses": reused_images, "imageBatches": image_batches,
              "batchImageLimit": args.max_unique_images, "allBatches": args.all_batches, "externalCalls": 0,
              "qualityAccepted": False, "extractionStatusChanged": False}
    (output/"review.json").write_text(json.dumps(review, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"reviewObjects": len(records), "uniqueImagesProcessed": review["uniqueImagesProcessed"], "imageBatches": image_batches, "externalCalls": 0,
                      "qualityAccepted": False, "extractionStatusChanged": False}))


if __name__ == "__main__":
    main()
