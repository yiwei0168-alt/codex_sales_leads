from __future__ import annotations
from pathlib import Path
from hashlib import sha256
import json
import posixpath
import zipfile
import xml.etree.ElementTree as ET
import pymupdf
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

EXTRACTOR_VERSION = "layout-v2.0.10"

def _clean(value):
    return " ".join(str("" if value is None else value).replace("\x00", "").split()).strip()

def _block(block_id, unit_type, unit_index, block_type, text="", quality="success", **extra):
    return {"id": block_id, "unitType": unit_type, "unitIndex": unit_index, "blockType": block_type,
            "text": text, "extractorVersion": EXTRACTOR_VERSION, "quality": quality, **extra}

def extract_pdf(path: Path):
    blocks=[]; statuses=[]
    for index,page in enumerate(pymupdf.open(path),1):
        try:
            page_blocks=[]
            for block_index,item in enumerate(page.get_text("blocks",sort=True),1):
                x0,y0,x1,y1,text,*_=item
                text=text.strip()
                if text: page_blocks.append(_block(f"page-{index}-block-{block_index}","page",index,"paragraph",text,bbox=[x0,y0,x1,y1]))
            if page_blocks:
                blocks.extend(page_blocks)
                visual_content=bool(page.get_images() or page.get_drawings() or list(page.annots() or []) or list(page.widgets() or []))
                invisible_text=not visual_content and any(trace.get("chars") and (trace.get("type")==3 or trace.get("opacity",1)<=0.01)
                    for trace in page.get_texttrace())
                if visual_content or invisible_text:
                    blocks.append(_block(f"page-{index}-visual-review","page",index,"pending-ocr","",quality="pending-ocr",
                        bbox=list(page.rect),metadata={"reviewReason":"invisible-pdf-text" if invisible_text else "mixed-page-visual-content"}))
                    statuses.append("pending-ocr")
                else:
                    statuses.append("success")
            elif (all(not page.parent.xref_stream(ref).strip() for ref in page.get_contents())
                  and not page.get_images() and not list(page.annots() or [])
                  and not list(page.widgets() or [])):
                # Only certify structurally empty pages. Image/painted/annotated
                # pages still need OCR/review, even if text extraction is empty.
                blocks.append(_block(f"page-{index}-blank","page",index,"blank","",quality="blank"))
                statuses.append("blank")
            else:
                blocks.append(_block(f"page-{index}-ocr","page",index,"pending-ocr","",quality="pending-ocr"))
                statuses.append("pending-ocr")
        except Exception as error:
            blocks.append(_block(f"page-{index}-failed","page",index,"note",type(error).__name__,quality="failed"))
            statuses.append("failed")
    return blocks,statuses

def extract_pptx(path: Path):
    blocks=[];statuses=[]
    def has_native_line(shapes):
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP and has_native_line(shape.shapes):
                return True
            element=shape._element
            if element.tag.endswith("}cxnSp") or any(node.get("prst")=="line" for node in element.iter() if node.tag.endswith("}prstGeom")):
                return True
        return False
    for index,slide in enumerate(Presentation(path).slides,1):
        start=len(blocks)
        # A table can hold diagram labels while separate native lines encode
        # relationships. Cell adjacency alone must not certify those claims.
        table_with_lines=has_native_line(slide.shapes)
        def visit(shapes, ordinals=(), shape_ids=()):
            for ordinal,shape in enumerate(shapes,1):
                shape_path=(*ordinals,ordinal);id_path=(*shape_ids,shape.shape_id)
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    visit(shape.shapes,shape_path,id_path)
                    continue
                key="-".join(map(str,shape_path))
                location={"shapePath":list(shape_path),"shapeIdPath":list(id_path)}
                if getattr(shape,"has_table",False):
                    rows=[[_clean(cell.text) for cell in row.cells] for row in shape.table.rows]
                    if rows:
                        blocks.append(_block(f"slide-{index}-table-{key}","slide",index,"table",
                            quality="pending-ocr" if table_with_lines else "success",
                            table={"headers":rows[0],"rows":rows[1:],"startRow":2},
                            **({"metadata":{"reviewReason":"table-with-native-lines"}} if table_with_lines else {}),**location))
                elif getattr(shape,"has_text_frame",False):
                    text="\n".join(_clean(p.text) for p in shape.text_frame.paragraphs if _clean(p.text))
                    if text: blocks.append(_block(f"slide-{index}-text-{key}","slide",index,"paragraph",text,**location))
                else:
                    # An unparsed visual object is not a blank slide. Retain its
                    # locator and prevent publication until local review resolves it.
                    blocks.append(_block(f"slide-{index}-visual-{key}","slide",index,"pending-ocr","",quality="pending-ocr",
                        metadata={"reviewReason":"unparsed-slide-visual","shapeType":str(shape.shape_type)},**location))
        visit(slide.shapes)
        pending=any(block["quality"]=="pending-ocr" for block in blocks[start:])
        statuses.append("pending-ocr" if pending else "success" if len(blocks)>start else "blank")
    return blocks,statuses

def xlsx_visual_references(path: Path):
    # Read source XML: openpyxl can discard unsupported drawings during load.
    ns={"s":"http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "r":"http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
    result={}
    with zipfile.ZipFile(path) as archive:
        relations={item.get("Id"):item for item in ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))}
        for sheet in ET.fromstring(archive.read("xl/workbook.xml")).findall("s:sheets/s:sheet",ns):
            relation=relations[sheet.get("{"+ns["r"]+"}id")]
            if relation.get("TargetMode")=="External":
                raise ValueError("External worksheet relationship")
            target=relation.get("Target", "")
            part=posixpath.normpath(target.lstrip("/") if target.startswith("/") else posixpath.join("xl",target))
            if not part.startswith("xl/"): raise ValueError("Invalid worksheet relationship")
            refs=[]
            for element in ET.fromstring(archive.read(part)).iter():
                kind=element.tag.rsplit("}",1)[-1]
                if kind in {"drawing","legacyDrawing","legacyDrawingHF","picture","oleObjects","controls"}:
                    refs.append({"kind":kind,"relationshipId":element.get("{"+ns["r"]+"}id"),"worksheetPart":part})
            result[sheet.get("name")]=refs
    return result


def xlsx_merged_ranges(path: Path):
    # Read merge coordinates from source XML; openpyxl's read-only mode does
    # not expose merged_cells but skips unsupported table-style metadata.
    ns={"s":"http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "r":"http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
    result={}
    with zipfile.ZipFile(path) as archive:
        relations={item.get("Id"):item for item in ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))}
        for sheet in ET.fromstring(archive.read("xl/workbook.xml")).findall("s:sheets/s:sheet",ns):
            relation=relations[sheet.get("{"+ns["r"]+"}id")]
            if relation.get("TargetMode")=="External":
                raise ValueError("External worksheet relationship")
            target=relation.get("Target", "")
            part=posixpath.normpath(target.lstrip("/") if target.startswith("/") else posixpath.join("xl",target))
            if not part.startswith("xl/"): raise ValueError("Invalid worksheet relationship")
            root=ET.fromstring(archive.read(part))
            result[sheet.get("name")]=[item.get("ref") for item in root.findall("s:mergeCells/s:mergeCell",ns)]
    return result


def extract_xlsx(path: Path):
    visuals=xlsx_visual_references(path)
    source_merges=xlsx_merged_ranges(path)
    workbook=load_workbook(path,read_only=True,data_only=True)
    blocks=[];statuses=[]
    for index,name in enumerate(workbook.sheetnames,1):
        sheet=workbook[name]
        start=len(blocks)
        rows=[];cached_errors=[]
        if hasattr(sheet,"iter_rows"):
            for cell_row in sheet.iter_rows(values_only=False):
                rows.append([_clean(cell.value) for cell in cell_row])
                for cell in cell_row:
                    if cell.data_type=="e": cached_errors.append({"cell":cell.coordinate,"value":_clean(cell.value)})
        merged=source_merges.get(sheet.title,[])
        # Keep physical row positions, including leading/interior blank rows.
        # Removing them makes startRow=2 point at different original cells.
        if any(any(row) for row in rows):
            width=max(len(row) for row in rows);headers=(rows[0]+[""]*width)[:width]
            error_note=f"Cached spreadsheet errors ({len(cached_errors)}): "+", ".join(f"{item['cell']} {item['value']}" for item in cached_errors[:32]) if cached_errors else ""
            metadata={"mergedRanges":merged}
            if cached_errors: metadata.update({"reviewReason":"cached-spreadsheet-error","cachedErrorCount":len(cached_errors),"cachedErrorCells":cached_errors[:32]})
            blocks.append(_block(f"sheet-{index}-table","sheet",index,"table",text=error_note,
                quality="pending-ocr" if cached_errors else "success",section=sheet.title,
                table={"headers":headers,"rows":[(row+[""]*width)[:width] for row in rows[1:]],"startRow":2,"headerRow":1,"startColumn":1,"cellRange":f"A1:{get_column_letter(width)}{len(rows)}"},metadata=metadata))
        for ordinal,reference in enumerate(visuals.get(sheet.title,[]),1):
            blocks.append(_block(f"sheet-{index}-visual-{ordinal}","sheet",index,"pending-ocr","",quality="pending-ocr",
                section=sheet.title,visualReference=reference,metadata={"reviewReason":"unparsed-sheet-visual","sourceReference":reference}))
        statuses.append("pending-ocr" if visuals.get(sheet.title) or cached_errors else "success" if len(blocks)>start else "blank")
    workbook.close()
    return blocks,statuses

def extract_file(path: Path, storage_key: str):
    suffix=path.suffix.lower()
    if suffix==".pdf": blocks,statuses=extract_pdf(path)
    elif suffix==".pptx": blocks,statuses=extract_pptx(path)
    elif suffix==".xlsx": blocks,statuses=extract_xlsx(path)
    elif suffix in {".txt",".md"}:
        text=path.read_text(encoding="utf-8");blocks=[_block("document-text","document",1,"paragraph",text)];statuses=["success" if text.strip() else "blank"]
    else: raise ValueError(f"Unsupported controlled asset: {suffix}")
    return {"storageKey":storage_key,"sourceSha256":sha256(path.read_bytes()).hexdigest(),"extractorVersion":EXTRACTOR_VERSION,"blocks":blocks,
            "qualitySummary":{key:statuses.count(key) for key in ("success","blank","pending-ocr","failed")}}

def write_artifact(path: Path, documents):
    payload={"extractorVersion":EXTRACTOR_VERSION,"documents":documents};path.parent.mkdir(parents=True,exist_ok=True)
    path.write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding="utf-8")
    return sha256(path.read_bytes()).hexdigest()
